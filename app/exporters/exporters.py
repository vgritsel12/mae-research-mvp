from __future__ import annotations

import csv
import io
import json
import uuid
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import OUTPUT_DIR, Settings, get_settings
from app.domain.enums import RunMode
from app.domain.models import (
    Article,
    BaselineScore,
    BaselineSnapshot,
    ChangeLog,
    EvidenceObservation,
    MatrixOverrideLog,
    MatrixScore,
    ResearchView,
    ReviewRun,
    ScenarioAssessment,
    ScenarioCard,
    ShiftSignal,
    Source,
    UpdateJob,
)
from app.services.baseline import active_baseline_snapshot, current_mae_cells, mae_status_metrics
from app.services.normalization import (
    CANONICAL_REGIONS,
    canonical_cell_for,
    mae_result_template_row_keys,
    parse_template_rows,
    template_region,
    valid_score,
    xlsx_general_assessment_row_key,
)
from app.services.production import (
    ARCHIVED_TEST_DATA,
    INVALID_REVIEW_STATUSES,
    INVALID_SIGNAL_STATUSES,
    has_test_marker,
    is_archived_test_record,
    is_content_valid_article,
    is_production_assessment,
    is_production_change,
    is_production_matrix_score,
    is_production_research_view,
    is_production_scenario,
    is_production_signal,
    is_production_source,
    production_scenario_ids,
)
from app.services.production_cleanup import invalidated_legacy_register


REGIONS = CANONICAL_REGIONS


class ExportService:
    def __init__(
        self,
        settings: Settings | None = None,
        export_mode: str | RunMode = RunMode.REAL.value,
        run_id: str | None = None,
    ) -> None:
        self.settings = settings or get_settings()
        self.export_mode = RunMode.REAL.value
        self.run_id = run_id or str(uuid.uuid4())
        self.exported_at = datetime.now(UTC)
        self._row_cache: dict[str, list[Any]] = {}
        self._filter_context_cache: dict[str, set[str]] | None = None

    def export_csv_bytes(self, session: Session) -> bytes:
        matrix = self._matrix_lookup(session)
        rows = [
            ["", "", "", "Geography", "Geography", "Geography", "Geography", "Geography", "Geography"],
            ["Asset Class", "Group", "Segment", *REGIONS],
        ]
        for row in parse_template_rows(self.settings.mae_template_xlsx):
            values = [row.asset_class, row.asset_group, row.asset_segment]
            for region in REGIONS:
                cell = canonical_cell_for(row.template_row_key, region, include_not_applicable=True)
                if cell is None or cell.applicability == "NOT_APPLICABLE":
                    values.append("")
                    continue
                score = matrix.get((row.template_row_key, region))
                values.append("" if score is None else str(valid_score(score)))
            rows.append(values)
        out = io.StringIO()
        writer = csv.writer(out, lineterminator="\n")
        writer.writerows(rows)
        return out.getvalue().encode("utf-8-sig")

    def export_json_bytes(self, session: Session) -> bytes:
        payload = {
            "schema_version": "1.0",
            "run": self._run_metadata(session),
            "exported_at": self.exported_at.isoformat(),
            "sources": [self._model_dict(x) for x in self._rows(session, Source)],
            "articles": [self._model_dict(x) for x in self._rows(session, Article)],
            "research_views": [self._model_dict(x) for x in self._rows(session, ResearchView)],
            "change_logs": [self._model_dict(x) for x in self._rows(session, ChangeLog)],
            "scenarios": [self._model_dict(x) for x in self._rows(session, ScenarioCard)],
            "evidence": [self._model_dict(x) for x in self._rows(session, EvidenceObservation)],
            "assessments": [self._model_dict(x) for x in self._rows(session, ScenarioAssessment)],
            "shift_signals": [self._model_dict(x) for x in self._rows(session, ShiftSignal)],
            "matrix_scores": [self._model_dict(x) for x in self._rows(session, MatrixScore)],
            "baseline": [self._model_dict(x) for x in self._rows(session, BaselineScore)],
            "current_mae": [self._current_cell_dict(x) for x in current_mae_cells(session, self.export_mode)],
            "financial_reviews": [self._model_dict(x) for x in self._rows(session, ReviewRun)],
            "manual_overrides": [self._model_dict(x) for x in self._rows(session, MatrixOverrideLog)],
            "invalidated_legacy_register": invalidated_legacy_register(session),
        }
        return json.dumps(payload, ensure_ascii=False, indent=2, default=self._json_default).encode("utf-8")

    def export_xlsx_bytes(self, session: Session) -> bytes:
        wb = load_workbook(self.settings.mae_template_xlsx)
        for sheet_name in [
            "Current MAE Matrix",
            "Legacy Template Compatibility",
            "MAE Matrix",
            "System Matrix",
            "Signals",
            "Scenarios",
            "Evidence",
            "Audit",
            "Articles",
            "Research Views",
            "Changes",
            "Baseline Matrix",
            "Validated Adjustments",
            "Current MAE",
            "Fresh Coverage",
            "Source Register",
            "Publication Register",
            "Evidence Register",
            "Confidence",
            "Automated Quality Diagnostics",
            "Score Adjustment Log",
            "Invalidated Legacy Register",
            "Report Metadata",
            "MAE Result",
        ]:
            if sheet_name in wb.sheetnames:
                del wb[sheet_name]
        if "ex" in wb.sheetnames:
            legacy = wb.copy_worksheet(wb["ex"])
            legacy.title = "Legacy Template Compatibility"
            legacy["A1"] = "LEGACY / DEPRECATED - portfolio allocation blocks are preserved only for compatibility"
            wb["ex"].title = "Current MAE Matrix"
            self._write_template_result_sheet(wb["Current MAE Matrix"], session)
        self._add_matrix_sheet(wb, session, sheet_name="System Matrix")
        self._add_table_sheet(
            wb,
            "Signals",
            ["id", "canonical_cell_id", "asset", "row_key", "region", "direction", "strength", "evidence", "pricing", "status", "sources"],
            [
                [
                    s.id,
                    s.canonical_cell_id,
                    s.asset,
                    s.template_row_key,
                    s.region,
                    s.direction,
                    s.suggested_strength,
                    s.evidence_status,
                    s.pricing_status,
                    s.analyst_status,
                    "\n".join(s.source_urls),
                ]
                for s in self._rows(session, ShiftSignal)
            ],
        )
        self._add_table_sheet(
            wb,
            "Scenarios",
            ["id", "change_id", "view_id", "type", "title", "triggers", "reversal", "expected_reaction"],
            [
                [
                    s.id,
                    s.linked_change_id,
                    s.linked_view_id,
                    s.scenario_type,
                    s.title,
                    "\n".join(s.triggers),
                    "\n".join(s.reversal_conditions),
                    s.expected_reaction,
                ]
                for s in self._rows(session, ScenarioCard)
            ],
        )
        self._add_table_sheet(
            wb,
            "Evidence",
            ["scenario_id", "indicator", "date", "expected", "actual", "importance", "support", "source_url"],
            [
                [
                    e.scenario_id,
                    e.indicator,
                    e.observation_date.isoformat(),
                    e.expected,
                    e.actual,
                    e.importance,
                    e.support_value,
                    e.source_url,
                ]
                for e in self._rows(session, EvidenceObservation)
            ],
        )
        self._add_table_sheet(
            wb,
            "Audit",
            ["row_key", "region", "suggested", "approved", "signals", "source_urls", "details"],
            [
                [
                    m.template_row_key,
                    m.region,
                    m.suggested_score,
                    m.approved_score,
                    "\n".join(m.signal_ids),
                    "\n".join(_signal_urls(session, m.signal_ids)),
                    json.dumps(m.calculation_details, ensure_ascii=False, default=self._json_default),
                ]
                for m in self._rows(session, MatrixScore)
                if m.signal_ids or m.approved_score is not None or m.override_reason
            ],
        )
        self._add_article_sheet(wb, session)
        self._add_research_views_sheet(wb, session)
        self._add_changes_sheet(wb, session)
        self._add_baseline_sheet(wb, session)
        self._add_validated_adjustments_sheet(wb, session)
        self._add_current_mae_sheet(wb, session)
        self._add_fresh_coverage_sheet(wb, session)
        self._add_register_sheets(wb, session)
        self._add_review_sheets(wb, session)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def write_exports(self, session: Session, output_dir: Path | None = None) -> dict[str, Path]:
        output_dir = output_dir or OUTPUT_DIR
        output_dir.mkdir(parents=True, exist_ok=True)
        paths = {
            "csv": output_dir / "mae_matrix_export.csv",
            "json": output_dir / "mae_export.json",
            "xlsx": output_dir / "mae_export.xlsx",
        }
        paths["csv"].write_bytes(self.export_csv_bytes(session))
        paths["json"].write_bytes(self.export_json_bytes(session))
        paths["xlsx"].write_bytes(self.export_xlsx_bytes(session))
        return paths

    def _matrix_lookup(self, session: Session) -> dict[tuple[str, str], int | None]:
        matrix: dict[tuple[str, str], int | None] = {}
        for row in current_mae_cells(session, self.export_mode):
            if row.current_score is None:
                continue
            matrix[(row.template_row_key, template_region(row.region))] = valid_score(row.current_score)
        return matrix

    def _add_matrix_sheet(self, wb, session: Session, sheet_name: str = "System Matrix") -> None:
        ws = wb.create_sheet(sheet_name, 1)
        matrix = self._matrix_lookup(session)
        ws.append(["Asset Class", "Geography", "Segment", *REGIONS])
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for row in parse_template_rows(self.settings.mae_template_xlsx):
            values = [
                row.asset_class,
                row.asset_group,
                row.asset_segment,
                *[matrix.get((row.template_row_key, region)) for region in REGIONS],
            ]
            ws.append(values)
        ws.freeze_panes = "D2"
        for col in ["A", "B", "C"]:
            ws.column_dimensions[col].width = 22
        for col in ["D", "E", "F", "G", "H", "I"]:
            ws.column_dimensions[col].width = 12
        for row in ws.iter_rows(min_row=2, min_col=4, max_col=9):
            for cell in row:
                if cell.value is not None:
                    valid_score(int(cell.value))
                cell.alignment = Alignment(horizontal="center")

    def _write_template_result_sheet(self, ws, session: Session) -> None:
        for merged in list(ws.merged_cells.ranges):
            if merged.min_col > 8 or merged.max_col > 8:
                ws.unmerge_cells(str(merged))
        if ws.max_column > 8:
            ws.delete_cols(9, ws.max_column - 8)
        for coordinate in list(ws._cells):
            if coordinate[1] > 8:
                del ws._cells[coordinate]
        for column_letter in list(ws.column_dimensions):
            if column_index_from_string(column_letter) > 8:
                del ws.column_dimensions[column_letter]
        matrix = self._matrix_lookup(session)
        # General Assessment area in the provided XLSX template is A:H. We
        # write Current MAE scores: baseline plus validated fresh adjustment.
        for merged in list(ws.merged_cells.ranges):
            if merged.min_col >= 3 and merged.max_col <= 8 and merged.min_row >= 3 and merged.max_row <= 22:
                ws.unmerge_cells(str(merged))
        current_asset_class = ""
        for row in range(3, 22):
            if ws[f"A{row}"].value:
                current_asset_class = str(ws[f"A{row}"].value)
            label = str(ws[f"B{row}"].value or "").strip()
            row_key = xlsx_general_assessment_row_key(current_asset_class, label)
            for col, region in zip(["C", "D", "E", "F", "G", "H"], REGIONS):
                ws[f"{col}{row}"].value = matrix.get((row_key, region)) if row_key else None
                ws[f"{col}{row}"].alignment = Alignment(horizontal="center")
        stats = self._run_metadata(session)
        warning = self._export_warning(session)
        status = mae_status_metrics(session, self.export_mode)
        header_rows = [
            ("Run ID", stats["run_id"]),
            ("Run date", stats["exported_at"]),
            ("Baseline date", status["baseline_date"] or status["baseline_date_status"]),
            (
                "Coverage",
                (
                    f"raw {status['raw_coverage_pct']}%; "
                    f"qualified {status['qualified_coverage_pct']}%; "
                    f"scenario-ready {status['financial_reviewed_coverage_pct']}%"
                ),
            ),
            ("Scenario adjustments", status["approved_shifts"]),
            ("Signals", stats["included_signals"]),
            ("Warning", warning),
        ]
        ws["A25"] = "MAE export header"
        ws["A25"].font = Font(bold=True, color="FFFFFF")
        ws["A25"].fill = PatternFill("solid", fgColor="9C0006" if warning else "305496")
        for offset, (label, value) in enumerate(header_rows, start=2):
            target_row = 24 + offset
            ws[f"A{target_row}"] = label
            ws[f"B{target_row}"] = value
            ws[f"A{target_row}"].font = Font(bold=True)
            ws[f"B{target_row}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws["B32"].comment = Comment("Current MAE = imported baseline plus validated fresh adjustment.", "MAE MVP")

    def _add_table_sheet(self, wb, name: str, headers: list[str], rows: list[list[Any]]) -> None:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append([_sanitize_export_value(value) for value in row])
        ws.freeze_panes = "A2"
        for column_cells in ws.columns:
            length = max(len(str(cell.value or "")) for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 55)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    def _add_article_sheet(self, wb, session: Session) -> None:
        sources = {s.id: s for s in session.scalars(select(Source)).all()}
        self._add_table_sheet(
            wb,
            "Articles",
            ["id", "date", "title", "source", "url", "fetch_status", "processing_status"],
            [
                [
                    a.id,
                    a.publication_date.isoformat(),
                    a.title,
                    sources[a.source_id].institution_name if a.source_id in sources else a.source_id,
                    a.source_reference,
                    a.fetch_status,
                    a.processing_status,
                ]
                for a in self._rows(session, Article)
            ],
        )

    def _add_research_views_sheet(self, wb, session: Session) -> None:
        articles = {a.id: a for a in session.scalars(select(Article)).all()}
        self._add_table_sheet(
            wb,
            "Research Views",
            [
                "id",
                "article_id",
                "article_date",
                "institution",
                "canonical_cell_id",
                "region",
                "row_key",
                "direction",
                "score",
                "confidence",
                "extraction_method",
                "quote",
                "source_url",
            ],
            [
                [
                    v.id,
                    v.article_id,
                    articles[v.article_id].publication_date.isoformat() if v.article_id in articles else "",
                    v.institution,
                    v.canonical_cell_id,
                    v.region,
                    v.template_row_key,
                    v.direction,
                    v.position_score,
                    v.confidence,
                    v.extraction_method,
                    (v.evidence_quotes[0].get("quote") if v.evidence_quotes else ""),
                    (v.evidence_quotes[0].get("source_url") if v.evidence_quotes else ""),
                ]
                for v in self._rows(session, ResearchView)
            ],
        )

    def _add_changes_sheet(self, wb, session: Session) -> None:
        views = {v.id: v for v in session.scalars(select(ResearchView)).all()}
        articles = {a.id: a for a in session.scalars(select(Article)).all()}
        self._add_table_sheet(
            wb,
            "Changes",
            [
                "id",
                "previous_view_id",
                "previous_url",
                "current_view_id",
                "current_url",
                "idea_state",
                "materiality",
                "change_types",
                "explanation",
            ],
            [
                [
                    c.id,
                    c.previous_view_id,
                    _view_url(c.previous_view_id, views, articles),
                    c.current_view_id,
                    _view_url(c.current_view_id, views, articles),
                    c.idea_state,
                    c.materiality,
                    ", ".join(c.change_types),
                    c.explanation,
                ]
                for c in self._rows(session, ChangeLog)
            ],
        )

    def _add_baseline_sheet(self, wb, session: Session) -> None:
        snapshot = active_baseline_snapshot(session)
        self._add_table_sheet(
            wb,
            "Baseline Matrix",
            ["snapshot_id", "baseline_date", "canonical_cell_id", "row_key", "region", "baseline_score", "source_row", "source_column"],
            [
                [
                    row.snapshot_id,
                    snapshot.baseline_date.isoformat() if snapshot and snapshot.baseline_date else "UNKNOWN",
                    row.canonical_cell_id,
                    row.template_row_key,
                    row.region,
                    row.baseline_score,
                    row.source_row,
                    row.source_column,
                ]
                for row in self._rows(session, BaselineScore)
            ],
        )

    def _add_validated_adjustments_sheet(self, wb, session: Session) -> None:
        rows = [cell for cell in current_mae_cells(session, self.export_mode) if cell.validated_adjustment is not None]
        self._add_table_sheet(
            wb,
            "Validated Adjustments",
            ["canonical_cell_id", "row_key", "region", "baseline", "validated_adjustment", "current_score", "sources", "confidence", "freshness"],
            [
                [
                    row.canonical_cell_id,
                    row.template_row_key,
                    row.region,
                    row.baseline_score,
                    row.validated_adjustment,
                    row.current_score,
                    row.source_count,
                    row.confidence,
                    row.freshness_score,
                ]
                for row in rows
            ],
        )

    def _add_current_mae_sheet(self, wb, session: Session) -> None:
        self._add_table_sheet(
            wb,
            "Current MAE",
            [
                "canonical_cell_id",
                "row_key",
                "region",
                "applicability",
                "baseline",
                "proposed_adjustment",
                "validated_adjustment",
                "current_score",
                "coverage_status",
            ],
            [
                [
                    row.canonical_cell_id,
                    row.template_row_key,
                    row.region,
                    row.applicability,
                    row.baseline_score,
                    row.proposed_adjustment,
                    row.validated_adjustment,
                    row.current_score,
                    row.coverage_status,
                ]
                for row in current_mae_cells(session, self.export_mode)
            ],
        )

    def _add_fresh_coverage_sheet(self, wb, session: Session) -> None:
        self._add_table_sheet(
            wb,
            "Fresh Coverage",
            ["canonical_cell_id", "row_key", "region", "coverage_status", "source_count", "publication_ids", "publication_dates", "supporting", "contradicting"],
            [
                [
                    row.canonical_cell_id,
                    row.template_row_key,
                    row.region,
                    row.coverage_status,
                    row.source_count,
                    "\n".join(row.publication_ids),
                    "\n".join(row.publication_dates),
                    json.dumps(row.supporting_evidence, ensure_ascii=False, default=self._json_default),
                    json.dumps(row.contradicting_evidence, ensure_ascii=False, default=self._json_default),
                ]
                for row in current_mae_cells(session, self.export_mode)
                if row.coverage_status not in {"NO_DATA", "NOT_APPLICABLE"}
            ],
        )

    def _add_register_sheets(self, wb, session: Session) -> None:
        self._add_table_sheet(
            wb,
            "Source Register",
            ["id", "institution", "website", "category", "adapter", "active", "last_checked"],
            [
                [s.id, s.institution_name, s.website, s.category, s.adapter_type, s.active, s.last_checked_at]
                for s in self._rows(session, Source)
            ],
        )
        self._add_table_sheet(
            wb,
            "Publication Register",
            ["id", "date", "title", "source_id", "canonical_url", "fetch_status", "processing_status"],
            [
                [a.id, a.publication_date.isoformat(), a.title, a.source_id, a.canonical_url, a.fetch_status, a.processing_status]
                for a in self._rows(session, Article)
            ],
        )
        self._add_table_sheet(
            wb,
            "Evidence Register",
            ["id", "scenario_id", "indicator", "date", "actual", "support", "source_url"],
            [
                [e.id, e.scenario_id, e.indicator, e.observation_date.isoformat(), e.actual, e.support_value, e.source_url]
                for e in self._rows(session, EvidenceObservation)
            ],
        )
        self._add_table_sheet(
            wb,
            "Confidence",
            ["row_key", "region", "confidence", "freshness", "source_count", "coverage_status"],
            [
                [row.template_row_key, row.region, row.confidence, row.freshness_score, row.source_count, row.coverage_status]
                for row in current_mae_cells(session, self.export_mode)
            ],
        )

    def _add_review_sheets(self, wb, session: Session) -> None:
        self._add_table_sheet(
            wb,
            "Automated Quality Diagnostics",
            ["role", "outcome", "overall", "source_quality", "evidence_logic", "consistency", "usefulness", "critical", "cell_findings"],
            [
                [
                    r.role,
                    _trust_safe_report_text(r.verdict),
                    r.overall_quality,
                    r.source_quality,
                    r.evidence_to_score_logic,
                    r.cross_matrix_consistency,
                    r.practical_usefulness,
                    _trust_safe_report_text(json.dumps(r.critical_findings, ensure_ascii=False, default=self._json_default)),
                    _trust_safe_report_text(json.dumps(r.cell_findings, ensure_ascii=False, default=self._json_default)),
                ]
                for r in self._rows(session, ReviewRun)
            ],
        )
        self._add_table_sheet(
            wb,
            "Score Adjustment Log",
            ["row_key", "region", "previous_score", "new_score", "reason", "actor", "created_at"],
            [
                [
                    row.template_row_key,
                    row.region,
                    row.previous_score,
                    row.new_score,
                    row.reason,
                    row.reviewer_label,
                    row.created_at,
                ]
                for row in self._rows(session, MatrixOverrideLog)
            ],
        )
        self._add_table_sheet(
            wb,
            "Invalidated Legacy Register",
            ["type", "id", "title", "institution", "url", "status", "reason", "updated_at"],
            [
                [
                    row["type"],
                    row["id"],
                    row["title"],
                    row["institution"],
                    row["url"],
                    row["status"],
                    row["reason"],
                    row["updated_at"],
                ]
                for row in invalidated_legacy_register(session)
            ],
        )
        metadata = self._run_metadata(session)
        status = mae_status_metrics(session, self.export_mode)
        self._add_table_sheet(
            wb,
            "Report Metadata",
            ["key", "value"],
            [[key, value] for key, value in {**metadata, **status}.items()],
        )

    def _rows(self, session: Session, model: Any) -> list[Any]:
        cache_key = model.__name__
        if cache_key in self._row_cache:
            return self._row_cache[cache_key]
        rows = list(session.scalars(select(model)).all())
        context = self._filter_context(session)
        if model is Source:
            filtered = [row for row in rows if row.id in context["source_ids"]]
        elif model is Article:
            filtered = [row for row in rows if row.id in context["article_ids"]]
        elif model is ResearchView:
            filtered = [row for row in rows if row.id in context["view_ids"]]
        elif model is ChangeLog:
            filtered = [row for row in rows if row.id in context["change_ids"]]
        elif model is ScenarioCard:
            filtered = [row for row in rows if row.id in context["scenario_ids"]]
        elif model is EvidenceObservation:
            filtered = [row for row in rows if row.scenario_id in context["scenario_ids"] and not row.is_demo and not has_test_marker(row.source_url)]
        elif model is ScenarioAssessment:
            filtered = [row for row in rows if row.scenario_id in context["scenario_ids"] and not row.is_demo and row.evidence_status != ARCHIVED_TEST_DATA]
        elif model is ShiftSignal:
            filtered = [row for row in rows if row.id in context["signal_ids"]]
        elif model is MatrixScore:
            filtered = [row for row in rows if row.id in context["matrix_ids"]]
        elif model is MatrixOverrideLog:
            filtered = [row for row in rows if not row.is_demo and not has_test_marker(row.reason)]
        elif model is ReviewRun:
            filtered = _latest_review_rows(rows)
        else:
            filtered = rows
        self._row_cache[cache_key] = filtered
        return filtered

    def _filter_context(self, session: Session) -> dict[str, set[str]]:
        if self._filter_context_cache is not None:
            return self._filter_context_cache
        sources = list(session.scalars(select(Source)).all())
        source_ids = {row.id for row in sources if is_production_source(row)}
        articles = list(session.scalars(select(Article)).all())
        article_ids = {row.id for row in articles if row.source_id in source_ids and is_content_valid_article(row)}
        views = list(session.scalars(select(ResearchView)).all())
        articles_by_id = {row.id: row for row in articles}
        view_ids = {
            row.id
            for row in views
            if row.article_id in article_ids
            and row.review_status not in INVALID_REVIEW_STATUSES
            and is_production_research_view(row, articles_by_id.get(row.article_id))
        }
        changes = list(session.scalars(select(ChangeLog)).all())
        change_ids = {
            row.id
            for row in changes
            if not is_archived_test_record(row) and row.current_view_id in view_ids
        }
        scenarios = list(session.scalars(select(ScenarioCard)).all())
        scenario_ids = {
            row.id
            for row in scenarios
            if not is_archived_test_record(row)
            and row.review_status not in INVALID_REVIEW_STATUSES
            and ((row.linked_change_id and row.linked_change_id in change_ids) or (row.linked_view_id and row.linked_view_id in view_ids))
        }
        signals = list(session.scalars(select(ShiftSignal)).all())
        signal_ids = {
            row.id
            for row in signals
            if not is_archived_test_record(row)
            and row.analyst_status not in INVALID_SIGNAL_STATUSES
            and row.evidence_status != ARCHIVED_TEST_DATA
            and (not row.change_id or row.change_id in change_ids)
            and not any(has_test_marker(url) for url in row.source_urls or [])
        }
        matrix_rows = list(session.scalars(select(MatrixScore)).all())
        matrix_ids = {
            row.id
            for row in matrix_rows
            if not row.is_demo
            and row.coverage_status != ARCHIVED_TEST_DATA
            and (not row.signal_ids or bool(set(row.signal_ids) & signal_ids) or bool(row.override_reason))
        }
        self._filter_context_cache = {
            "source_ids": source_ids,
            "article_ids": article_ids,
            "view_ids": view_ids,
            "change_ids": change_ids,
            "scenario_ids": scenario_ids,
            "signal_ids": signal_ids,
            "matrix_ids": matrix_ids,
        }
        return self._filter_context_cache

    def _run_metadata(self, session: Session) -> dict[str, Any]:
        mode_signals = self._rows(session, ShiftSignal)
        return {
            "run_id": self.run_id,
            "exported_at": self.exported_at.isoformat(),
            "environment": "production",
            "included_signals": len(mode_signals),
        }

    def _export_warning(self, session: Session) -> str:
        signals = self._rows(session, ShiftSignal)
        mapped_keys = mae_result_template_row_keys()
        unmapped = [s.template_row_key for s in signals if s.template_row_key not in mapped_keys]
        warnings = []
        if not [s for s in signals if s.analyst_status == "APPROVED"]:
            warnings.append("No approved mapped Shift Signals; Current MAE remains baseline unless manual overrides exist.")
        if unmapped:
            warnings.append("Unmapped signals excluded from Current MAE Matrix: " + ", ".join(sorted(set(unmapped))))
        return " ".join(warnings)

    @staticmethod
    def _model_dict(row: Any) -> dict[str, Any]:
        payload = {col.name: getattr(row, col.name) for col in row.__table__.columns}
        payload.pop("is_demo", None)
        return _sanitize_export_value(payload)

    @staticmethod
    def _current_cell_dict(row: Any) -> dict[str, Any]:
        return _sanitize_export_value({
            "canonical_cell_id": row.canonical_cell_id,
            "row_key": row.template_row_key,
            "region": row.region,
            "applicability": row.applicability,
            "baseline_score": row.baseline_score,
            "proposed_adjustment": row.proposed_adjustment,
            "validated_adjustment": row.validated_adjustment,
            "current_score": row.current_score,
            "coverage_status": row.coverage_status,
            "source_count": row.source_count,
            "publication_ids": row.publication_ids,
            "publication_dates": row.publication_dates,
            "supporting_evidence": row.supporting_evidence,
            "contradicting_evidence": row.contradicting_evidence,
            "confidence": row.confidence,
            "freshness_score": row.freshness_score,
        })

    @staticmethod
    def _json_default(value: Any) -> str:
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return str(value)


def _view_url(view_id: str | None, views: dict[str, ResearchView], articles: dict[str, Article]) -> str:
    if not view_id or view_id not in views:
        return ""
    article_id = views[view_id].article_id
    return articles[article_id].source_reference if article_id in articles else ""


def _signal_urls(session: Session, signal_ids: list[str]) -> list[str]:
    if not signal_ids:
        return []
    signals = session.scalars(select(ShiftSignal).where(ShiftSignal.id.in_(signal_ids))).all()
    urls: list[str] = []
    for signal in signals:
        for url in signal.source_urls:
            if url not in urls:
                urls.append(url)
    return urls


def _latest_review_rows(rows: list[ReviewRun]) -> list[ReviewRun]:
    latest: dict[str, ReviewRun] = {}
    for row in sorted(rows, key=_review_sort_key):
        metadata_text = json.dumps(row.run_metadata or {}, ensure_ascii=False, default=str)
        if has_test_marker(metadata_text):
            continue
        latest[row.role] = row
    return list(latest.values())


def _review_sort_key(row: ReviewRun) -> str:
    value = row.created_at
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(UTC).replace(tzinfo=None)
        return value.isoformat()
    return str(value)


def _trust_safe_report_text(value: Any) -> str:
    text = str(value or "")
    replacements = {
        "FINANCIAL PASS WITH NO APPROVED SHIFTS": "AUTOMATED QUALITY PASS WITH NO SCENARIO ADJUSTMENTS",
        "FINANCIAL PASS": "AUTOMATED QUALITY PASS",
        "FINANCIAL FAIL": "AUTOMATED QUALITY FAIL",
        "financial review": "automated quality diagnostic",
        "Financial review": "Automated quality diagnostic",
        "Financial Review": "Automated Quality Diagnostic",
        "financially approved": "automatically assessed",
        "Approved shifts": "Scenario adjustments",
        "approved shifts": "scenario adjustments",
        "reviewer_verdict": "diagnostic_outcome",
        "reviewer_comment": "diagnostic_comment",
        "reviewer": "diagnostic",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def _sanitize_export_value(value: Any) -> Any:
    if isinstance(value, str):
        return (
            value.replace("DEMO/RULE", "RULE")
            .replace("DEMO:", "TEST:")
            .replace("DEMO ", "TEST ")
            .replace("DEMO-", "TEST-")
            .replace("DEMO_", "TEST_")
        )
    if isinstance(value, list):
        return [_sanitize_export_value(item) for item in value]
    if isinstance(value, tuple):
        return tuple(_sanitize_export_value(item) for item in value)
    if isinstance(value, dict):
        return {str(key): _sanitize_export_value(item) for key, item in value.items()}
    return value
