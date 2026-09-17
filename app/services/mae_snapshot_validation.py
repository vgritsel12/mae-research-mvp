from __future__ import annotations

import json
import re
import sqlite3
import hashlib
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook

from app.config import ROOT_DIR, get_settings
from app.services.mae_snapshot_product import (
    BASELINE_DATE,
    BASELINE_EVIDENCE,
    BASELINE_SCORES,
    CHANGE_STATUSES,
    EVIDENCE_STATUSES,
    PRODUCTION_STATUS,
    PRODUCT_VERSION,
    FULL_ANCHORS,
    FULL_ASSETS,
    FULL_GEOGRAPHIES,
    FULL_PRODUCT_VERSION,
    REVIEW_STATUSES,
    TEST_ONLY_STATUS,
    database_state,
    read_csv,
    snapshot_status_for_date,
)


EXPECTED_SHEETS = [
    "Strategy_Sentiment_Map",
    "Change_Tracker",
    "Scenarios",
    "Evidence_Expectation",
    "Transmission_Control",
]
FORMULA_ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A")


def validate_snapshot(
    snapshot_date: date,
    *,
    root: Path = ROOT_DIR,
    snapshot_dir: Path | None = None,
    build_metadata: dict[str, Any] | None = None,
    require_database: bool = True,
    write_report: bool = True,
    allow_future_test_date: bool = False,
    today: date | None = None,
) -> dict[str, Any]:
    stamp = snapshot_date.isoformat()
    snapshot_dir = snapshot_dir or root / "outputs" / "snapshots" / stamp
    score_path = snapshot_dir / f"mae_scores_{stamp}.csv"
    evidence_path = snapshot_dir / f"mae_evidence_{stamp}.csv"
    scenario_path = snapshot_dir / f"mae_scenarios_{stamp}.csv"
    report_path = snapshot_dir / f"mae_report_{stamp}.md"
    workbook_path = snapshot_dir / f"mae_snapshot_{stamp}.xlsx"
    validation_path = snapshot_dir / f"validation_report_{stamp}.json"
    if build_metadata is None and validation_path.exists():
        try:
            build_metadata = json.loads(validation_path.read_text(encoding="utf-8")).get("metadata", {})
        except (OSError, json.JSONDecodeError):
            build_metadata = {}
    build_metadata = build_metadata or {}
    current_date = today or date.today()
    try:
        snapshot_status = snapshot_status_for_date(
            snapshot_date,
            allow_future_test_date=allow_future_test_date,
            today=current_date,
        )
    except ValueError:
        snapshot_status = PRODUCTION_STATUS
    build_metadata = dict(build_metadata)
    if snapshot_date > current_date and allow_future_test_date:
        snapshot_status = TEST_ONLY_STATUS
        build_metadata["snapshot_status"] = TEST_ONLY_STATUS
        build_metadata["production_eligible"] = False
    else:
        snapshot_status = str(build_metadata.get("snapshot_status") or snapshot_status)

    freshness: dict[str, Any] = {
        "snapshot_date": snapshot_date.isoformat(),
        "latest_publication_date": "",
        "latest_market_observation_date": str(
            (build_metadata.get("data_windows") or {}).get("latest_market_observation_date") or ""
        ),
        "previous_comparable_snapshot": "",
        "new_publications_since_previous": 0,
        "new_market_observations_since_previous": int(
            (build_metadata.get("data_windows") or {}).get("new_market_observations_since_previous") or 0
        ),
        "carry_forward_cells": 0,
        "changed_cells": 0,
        "warnings": [],
    }

    checks: list[dict[str, str]] = []
    errors: list[str] = []

    def check(name: str, condition: bool, detail: str) -> None:
        status = "PASS" if condition else "FAIL"
        checks.append({"name": name, "status": status, "detail": detail})
        if not condition:
            errors.append(f"{name}: {detail}")

    snapshot_date_ok = snapshot_date <= current_date or allow_future_test_date
    check(
        "snapshot_date_not_future",
        snapshot_date_ok,
        (
            f"snapshot_date={snapshot_date.isoformat()}, current_date={current_date.isoformat()}"
            if snapshot_date_ok
            else "Snapshot date cannot be in the future."
        ),
    )
    if not snapshot_date_ok:
        return _finish_report(
            snapshot_date,
            validation_path,
            checks,
            errors,
            build_metadata,
            write_report,
            freshness=freshness,
            snapshot_status=snapshot_status,
        )

    required = [score_path, evidence_path, scenario_path, report_path, workbook_path]
    missing = [str(path) for path in required if not path.exists()]
    check("required_files", not missing, "all required files exist" if not missing else f"missing: {', '.join(missing)}")
    if missing:
        return _finish_report(
            snapshot_date,
            validation_path,
            checks,
            errors,
            build_metadata,
            write_report,
            freshness=freshness,
            snapshot_status=snapshot_status,
        )

    scores = read_csv(score_path)
    evidence = read_csv(evidence_path)
    scenarios = read_csv(scenario_path)
    report_text = report_path.read_text(encoding="utf-8")

    cell_ids = [row.get("cell_id", "") for row in scores]
    check("ten_unique_cells", len(scores) == 10 and len(set(cell_ids)) == 10, f"rows={len(scores)}, unique={len(set(cell_ids))}")
    check("no_duplicate_cell_id", len(cell_ids) == len(set(cell_ids)), "cell_id values are unique")
    parsed_scores: dict[str, int] = {}
    score_error = ""
    for row in scores:
        try:
            value = int(row["score"])
            if not -3 <= value <= 3:
                raise ValueError
            parsed_scores[row["cell_id"]] = value
        except (KeyError, TypeError, ValueError):
            score_error = f"invalid score for {row.get('cell_id', '<missing>')}: {row.get('score')}"
            break
    check("valid_scores", not score_error and len(parsed_scores) == 10, score_error or "all scores are integers in -3..+3")
    check(
        "valid_change_statuses",
        all(row.get("change_status") in CHANGE_STATUSES for row in scores),
        "all change statuses are allowed",
    )

    pass_cells = {row.get("cell_id", "") for row in evidence if row.get("review_status") == "PASS"}
    check("pass_evidence_per_cell", set(cell_ids) <= pass_cells, f"supported={len(set(cell_ids) & pass_cells)}/10")
    check("coverage_10_of_10", len(set(cell_ids) & pass_cells) == 10, f"coverage={len(set(cell_ids) & pass_cells)}/10")
    check(
        "evidence_status_values",
        all(row.get("evidence_status") in EVIDENCE_STATUSES for row in evidence),
        "evidence statuses are CONFIRMED, MIXED or CONTRADICTED",
    )
    check(
        "review_status_values",
        all(row.get("review_status") in REVIEW_STATUSES for row in evidence),
        "review statuses are PASS, MANUAL_REVIEW or REJECTED",
    )
    date_errors = []
    future_publications = []
    pass_publication_dates: list[date] = []
    window_start = snapshot_date - timedelta(days=90)
    for row in evidence:
        try:
            publication_date = date.fromisoformat(row["publication_date"])
        except (KeyError, ValueError):
            date_errors.append(f"{row.get('source_id')}: invalid publication_date")
            continue
        if row.get("review_status") == "PASS":
            pass_publication_dates.append(publication_date)
        if publication_date > snapshot_date:
            future_publications.append(f"{row.get('source_id')}: {publication_date.isoformat()}")
        carry = str(row.get("carry_forward_flag", "")).lower() == "true"
        if publication_date < window_start and not carry:
            date_errors.append(f"{row.get('source_id')}: outside 90-day window without carry-forward")
    check(
        "publication_not_future",
        not future_publications,
        "all publication dates are on or before snapshot date"
        if not future_publications
        else f"future publication rejected: {'; '.join(future_publications[:8])}",
    )
    check(
        "evidence_dates",
        not date_errors,
        "90-day/carry-forward rule passed" if not date_errors else "; ".join(date_errors[:8]),
    )
    freshness["latest_publication_date"] = (
        max(pass_publication_dates).isoformat() if pass_publication_dates else ""
    )

    check("three_scenarios", len(scenarios) == 3, f"scenario rows={len(scenarios)}")
    check("scenario_names", {row.get("scenario_id") for row in scenarios} == {"BASE", "UPSIDE", "DOWNSIDE"}, "BASE, UPSIDE and DOWNSIDE are present")
    for field in ("narrative", "causal_chain", "trigger", "veto", "affected_cells"):
        check(f"scenario_{field}", all(str(row.get(field) or "").strip() for row in scenarios), f"all scenarios have {field}")

    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    check("five_excel_sheets", workbook.sheetnames == EXPECTED_SHEETS, f"sheets={workbook.sheetnames}")
    excel_scores = _read_table_map(workbook["Strategy_Sentiment_Map"], "asset", "score")
    csv_scores_by_asset = {row["segment"]: int(row["score"]) for row in scores}
    check("scores_excel_csv", excel_scores == csv_scores_by_asset, f"Excel scores={len(excel_scores)}, CSV scores={len(csv_scores_by_asset)}")
    markdown_scores = _markdown_scores(report_text)
    check("scores_markdown_csv", markdown_scores == parsed_scores, f"Markdown scores={len(markdown_scores)}, CSV scores={len(parsed_scores)}")

    excel_scenarios = _read_table_rows(workbook["Scenarios"], "scenario name")
    excel_scenario_map = {row.get("scenario name", ""): row.get("causal chain", "") for row in excel_scenarios}
    csv_scenario_map = {row["scenario_name"]: row["causal_chain"] for row in scenarios}
    check("scenarios_excel_csv", excel_scenario_map == csv_scenario_map, f"Excel scenarios={len(excel_scenario_map)}, CSV scenarios={len(csv_scenario_map)}")
    markdown_scenarios_ok = all(row["scenario_name"] in report_text and row["causal_chain"] in report_text for row in scenarios)
    check("scenarios_markdown_csv", markdown_scenarios_ok, "scenario names and causal chains match Markdown")

    formula_errors = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(token in value for token in FORMULA_ERRORS):
                    formula_errors.append(f"{worksheet.title}!{cell.coordinate}={value}")
    workbook.close()
    data_workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    for worksheet in data_workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value.startswith("#")):
                    formula_errors.append(f"{worksheet.title}!{cell.coordinate}={cell.value}")
    data_workbook.close()
    check("no_formula_errors", not formula_errors, "no #REF!, #VALUE!, #DIV/0!, #NAME? or #N/A" if not formula_errors else "; ".join(formula_errors[:8]))

    previous_dates = {row.get("previous_snapshot_date") for row in scores}
    previous_date = date.fromisoformat(next(iter(previous_dates))) if len(previous_dates) == 1 else None
    previous_ok = False
    if previous_date == BASELINE_DATE:
        previous_ok = (root / BASELINE_SCORES).exists() and (root / BASELINE_EVIDENCE).exists()
    elif previous_date:
        previous_dir = root / "outputs" / "snapshots" / previous_date.isoformat()
        previous_ok = (
            (previous_dir / f"mae_scores_{previous_date.isoformat()}.csv").exists()
            and (previous_dir / f"validation_report_{previous_date.isoformat()}.json").exists()
        )
    check("previous_snapshot_preserved", previous_ok, f"previous_snapshot={previous_date}")
    freshness["previous_comparable_snapshot"] = previous_date.isoformat() if previous_date else ""
    freshness["carry_forward_cells"] = sum(
        str(row.get("carry_forward") or "").lower() == "true" for row in scores
    )
    freshness["changed_cells"] = sum(int(row.get("score_delta") or 0) != 0 for row in scores)
    if previous_date:
        new_research = set()
        for row in evidence:
            if row.get("review_status") != "PASS" or "FRED" in str(row.get("provider") or "").upper():
                continue
            try:
                publication_date = date.fromisoformat(str(row.get("publication_date") or ""))
            except ValueError:
                continue
            if publication_date > previous_date:
                new_research.add(row.get("source_id"))
        freshness["new_publications_since_previous"] = len(new_research)
    if freshness["new_publications_since_previous"] == 0:
        freshness["warnings"].append("NO_NEW_RESEARCH_SINCE_PREVIOUS_SNAPSHOT")

    market_days = int((build_metadata.get("data_windows") or {}).get("market_window_days") or 0)
    macro_days = int((build_metadata.get("data_windows") or {}).get("macro_window_days") or 0)
    check("market_window_12m", market_days >= 365, f"market_window_days={market_days}")
    check("macro_window_24m", macro_days >= 730, f"macro_window_days={macro_days}")
    market_metadata = build_metadata.get("data_windows") or {}
    latest_market_text = str(market_metadata.get("latest_market_observation_date") or "")
    future_market_count = int(market_metadata.get("future_market_observation_count") or 0)
    try:
        latest_market_date = date.fromisoformat(latest_market_text) if latest_market_text else None
    except ValueError:
        latest_market_date = None
    market_date_ok = (
        latest_market_date is not None
        and latest_market_date <= snapshot_date
        and future_market_count == 0
    )
    check(
        "market_observation_not_future",
        market_date_ok,
        (
            f"latest_market_observation_date={latest_market_text}"
            if market_date_ok
            else (
                "Future market observation rejected: "
                f"latest={latest_market_text or 'missing'}, future_count={future_market_count}, as_of={snapshot_date.isoformat()}"
            )
        ),
    )
    if latest_market_date and latest_market_date < _latest_completed_trading_day(snapshot_date, current_date):
        freshness["warnings"].append("STALE_MARKET_DATA")

    if require_database:
        db_checks = _database_checks(snapshot_date, build_metadata, root=root)
        for name, condition, detail in db_checks:
            check(name, condition, detail)

    return _finish_report(
        snapshot_date,
        validation_path,
        checks,
        errors,
        build_metadata,
        write_report,
        freshness=freshness,
        snapshot_status=snapshot_status,
    )


def _database_checks(snapshot_date: date, metadata: dict[str, Any], *, root: Path) -> list[tuple[str, bool, str]]:
    current = database_state(root=root)
    preflight = metadata.get("database_preflight") or {}
    preflight_tables = preflight.get("tables") or {}
    monotonic = all(current.get("tables", {}).get(table, 0) >= count for table, count in preflight_tables.items())
    backups = list((root / "data" / "backups").glob("*.db"))
    snapshot_rows = cells = evidence = scenarios = 0
    database_path = Path(current["path"])
    if database_path.exists():
        with sqlite3.connect(f"file:{database_path}?mode=ro", uri=True) as connection:
            row = connection.execute(
                "SELECT id, run_metadata, status FROM mae_snapshots WHERE snapshot_date = ? AND is_demo = 0",
                (snapshot_date.isoformat(),),
            ).fetchone()
            if row:
                try:
                    run_metadata = json.loads(row[1]) if isinstance(row[1], str) else row[1]
                except (TypeError, json.JSONDecodeError):
                    run_metadata = {}
                if (
                    (run_metadata or {}).get("product_version") == PRODUCT_VERSION
                    and row[2] not in {"INVALID_FUTURE_DATE", TEST_ONLY_STATUS}
                ):
                    snapshot_rows = 1
                    cells = connection.execute("SELECT COUNT(*) FROM mae_snapshot_cells WHERE snapshot_id = ?", (row[0],)).fetchone()[0]
                    evidence = connection.execute("SELECT COUNT(*) FROM evidence_items WHERE snapshot_id = ? AND review_status = 'PASS'", (row[0],)).fetchone()[0]
            try:
                scenarios = connection.execute(
                    "SELECT COUNT(*) FROM market_scenarios WHERE snapshot_date = ? AND is_demo = 0",
                    (snapshot_date.isoformat(),),
                ).fetchone()[0]
            except sqlite3.OperationalError:
                scenarios = 0
    return [
        ("sqlite_backup_exists", bool(backups), f"backups={len(backups)}"),
        ("sqlite_history_not_reduced", monotonic, f"preflight={preflight_tables}, current={current.get('tables', {})}"),
        ("snapshot_persisted", snapshot_rows == 1 and cells >= 10, f"product snapshots={snapshot_rows}, cells={cells}, pilot anchors>=10"),
        ("pass_evidence_persisted", evidence >= 10, f"PASS evidence rows={evidence}"),
        ("scenarios_persisted", scenarios == 3, f"scenario rows={scenarios}"),
    ]


def _read_table_map(worksheet, key_header: str, value_header: str) -> dict[str, int]:
    rows = _read_table_rows(worksheet, key_header)
    result = {}
    for row in rows:
        try:
            result[str(row[key_header])] = int(row[value_header])
        except (KeyError, TypeError, ValueError):
            continue
    return result


def _read_table_rows(worksheet, first_header: str) -> list[dict[str, Any]]:
    values = list(worksheet.iter_rows(values_only=True))
    header_index = None
    headers: list[str] = []
    for index, row in enumerate(values):
        normalized = [str(value or "").strip().lower() for value in row]
        if first_header.lower() in normalized:
            header_index = index
            headers = normalized
            break
    if header_index is None:
        return []
    result = []
    for row in values[header_index + 1 :]:
        if not any(value is not None and str(value).strip() for value in row):
            continue
        record = {headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]}
        result.append(record)
    return result


def _markdown_scores(text: str) -> dict[str, int]:
    result = {}
    pattern = re.compile(r"^\|\s*([A-Z0-9_]+)\s*\|[^|]*\|\s*(-?\d+)\s*\|", re.MULTILINE)
    for match in pattern.finditer(text):
        result[match.group(1)] = int(match.group(2))
    return result


def _latest_completed_trading_day(snapshot_date: date, current_date: date) -> date:
    candidate = snapshot_date
    if snapshot_date >= current_date:
        candidate -= timedelta(days=1)
    while candidate.weekday() >= 5:
        candidate -= timedelta(days=1)
    return candidate


def _finish_report(
    snapshot_date: date,
    validation_path: Path,
    checks: list[dict[str, str]],
    errors: list[str],
    metadata: dict[str, Any],
    write_report: bool,
    *,
    freshness: dict[str, Any],
    snapshot_status: str,
) -> dict[str, Any]:
    report_status = "FAIL" if errors else TEST_ONLY_STATUS if snapshot_status == TEST_ONLY_STATUS else "PASS"
    report = {
        "status": report_status,
        "snapshot_status": snapshot_status,
        "snapshot_date": snapshot_date.isoformat(),
        "latest_publication_date": freshness.get("latest_publication_date", ""),
        "latest_market_observation_date": freshness.get("latest_market_observation_date", ""),
        "previous_comparable_snapshot": freshness.get("previous_comparable_snapshot", ""),
        "new_publications_since_previous": freshness.get("new_publications_since_previous", 0),
        "new_market_observations_since_previous": freshness.get("new_market_observations_since_previous", 0),
        "carry_forward_cells": freshness.get("carry_forward_cells", 0),
        "changed_cells": freshness.get("changed_cells", 0),
        "warnings": freshness.get("warnings", []),
        "checked_at": datetime.now(UTC).isoformat(),
        "coverage": next((row["detail"] for row in checks if row["name"] == "coverage_10_of_10"), "not checked"),
        "scenario_count": next((row["detail"] for row in checks if row["name"] == "three_scenarios"), "not checked"),
        "excel_sheet_count": next((row["detail"] for row in checks if row["name"] == "five_excel_sheets"), "not checked"),
        "checks": checks,
        "errors": errors,
        "metadata": metadata,
    }
    if write_report:
        validation_path.parent.mkdir(parents=True, exist_ok=True)
        validation_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def validate_full_snapshot(
    snapshot_date: date,
    *,
    root: Path = ROOT_DIR,
    snapshot_dir: Path | None = None,
    build_metadata: dict[str, Any] | None = None,
    require_database: bool = True,
    write_report: bool = True,
    allow_future_test_date: bool = False,
    today: date | None = None,
    pilot_output_paths: dict[str, Path] | None = None,
) -> dict[str, Any]:
    stamp = snapshot_date.isoformat()
    snapshot_dir = snapshot_dir or root / "outputs" / "snapshots" / stamp / "full"
    paths = {
        "scores": snapshot_dir / f"mae_full_scores_{stamp}.csv",
        "evidence": snapshot_dir / f"mae_full_evidence_{stamp}.csv",
        "scenarios": snapshot_dir / f"mae_full_scenarios_{stamp}.csv",
        "transmission": snapshot_dir / f"mae_full_transmission_{stamp}.csv",
        "report": snapshot_dir / f"mae_full_report_{stamp}.md",
        "quality_json": snapshot_dir / f"mae_full_quality_report_{stamp}.json",
        "quality_md": snapshot_dir / f"mae_full_quality_report_{stamp}.md",
        "workbook": snapshot_dir / f"mae_full_snapshot_{stamp}.xlsx",
        "validation": snapshot_dir / f"mae_full_validation_{stamp}.json",
    }
    if build_metadata is None and paths["validation"].exists():
        try:
            build_metadata = json.loads(paths["validation"].read_text(encoding="utf-8")).get("metadata", {})
        except (OSError, json.JSONDecodeError):
            build_metadata = {}
    metadata = dict(build_metadata or {})
    current_date = today or date.today()
    checks: list[dict[str, str]] = []
    errors: list[str] = []

    def check(name: str, condition: bool, detail: str) -> None:
        checks.append({"name": name, "status": "PASS" if condition else "FAIL", "detail": detail})
        if not condition:
            errors.append(f"{name}: {detail}")

    date_ok = snapshot_date <= current_date or allow_future_test_date
    check("snapshot_date_not_future", date_ok, snapshot_date.isoformat() if date_ok else "Snapshot date cannot be in the future.")
    if not date_ok:
        return _finish_full_validation(paths["validation"], metadata, checks, errors, write_report, snapshot_date)
    missing = [str(path) for key, path in paths.items() if key != "validation" and not path.exists()]
    check("required_files", not missing, "all full artifacts exist" if not missing else f"missing: {', '.join(missing)}")
    if missing:
        return _finish_full_validation(paths["validation"], metadata, checks, errors, write_report, snapshot_date)

    scores = read_csv(paths["scores"])
    evidence = read_csv(paths["evidence"])
    scenarios = read_csv(paths["scenarios"])
    transmission = read_csv(paths["transmission"])
    report_text = paths["report"].read_text(encoding="utf-8")
    quality_report = json.loads(paths["quality_json"].read_text(encoding="utf-8"))
    quality_markdown = paths["quality_md"].read_text(encoding="utf-8")
    cell_ids = [row.get("cell_id", "") for row in scores]
    applicable = [row for row in scores if row.get("applicable", "").lower() == "true"]
    not_applicable = [row for row in scores if row.get("applicable", "").lower() == "false"]
    expected_pairs = {(asset[2], geography) for asset in FULL_ASSETS for geography in FULL_GEOGRAPHIES}
    actual_pairs = {(row.get("asset_segment"), row.get("geography")) for row in scores}
    check("canonical_universe_114", len(scores) == 114 and len(set(cell_ids)) == 114 and actual_pairs == expected_pairs, f"rows={len(scores)}, unique={len(set(cell_ids))}")
    check("applicable_cells_104", len(applicable) == 104, f"applicable={len(applicable)}")
    expected_na = {
        (asset, geography)
        for asset in ("Gold", "Other Precious Metals")
        for geography in FULL_GEOGRAPHIES
        if geography != "Global"
    }
    actual_na = {(row["asset_segment"], row["geography"]) for row in not_applicable}
    check("not_applicable_cells_10", len(not_applicable) == 10 and actual_na == expected_na, f"N/A={len(not_applicable)}")

    parsed_scores: dict[str, int] = {}
    invalid_scores: list[str] = []
    for row in applicable:
        try:
            raw = row["score"]
            if not re.fullmatch(r"-?\d+", str(raw)):
                raise ValueError
            value = int(raw)
            if not -3 <= value <= 3:
                raise ValueError
            parsed_scores[row["cell_id"]] = value
        except (KeyError, TypeError, ValueError):
            invalid_scores.append(row.get("cell_id", "<missing>"))
    check("applicable_score_complete", not invalid_scores and len(parsed_scores) == 104, "0 missing scores" if not invalid_scores else f"invalid: {', '.join(invalid_scores[:20])}")
    check("score_integer_range", not invalid_scores, "all applicable scores are integers in -3..+3")
    na_score_ok = all(not str(row.get("score") or "").strip() for row in not_applicable)
    check("not_applicable_has_no_score", na_score_ok, "all N/A intersections have blank score")
    neutral_bad = [
        row["cell_id"]
        for row in applicable
        if int(row["score"]) == 0
        and row.get("evidence_mode") != "ANCHOR_EXISTING"
        and "сбалансирован" not in row.get("thesis", "").lower()
    ]
    check("zero_is_proven_neutral", not neutral_bad, "all zero scores have an explicit balanced thesis" if not neutral_bad else f"cells={', '.join(neutral_bad)}")

    pilot_scores = {row["cell_id"]: row for row in read_csv(root / "outputs" / "mae_latest_scores.csv")}
    full_scores = {row["cell_id"]: row for row in scores}
    anchor_errors = []
    for anchor_id in FULL_ANCHORS.values():
        full = full_scores.get(anchor_id, {})
        pilot = pilot_scores.get(anchor_id, {})
        if (
            full.get("score") != pilot.get("score")
            or full.get("thesis") != pilot.get("current_thesis")
            or full.get("evidence_status") != pilot.get("evidence_status_code")
            or full.get("evidence_mode") != "ANCHOR_EXISTING"
        ):
            anchor_errors.append(anchor_id)
    check("ten_anchors_preserved", not anchor_errors and len(FULL_ANCHORS) == 10, "10/10 immutable anchors preserved" if not anchor_errors else f"changed={', '.join(anchor_errors)}")

    review_bad = [row["cell_id"] for row in applicable if row.get("review_status") != "PASS"]
    check("all_applicable_review_pass", not review_bad, "104/104 assessment rows PASS" if not review_bad else f"cells={', '.join(review_bad[:20])}")
    benchmark_fields = {
        "instrument_definition", "benchmark_id", "benchmark_name", "benchmark_provider", "benchmark_type",
        "currency", "duration_bucket", "regional_scope", "market_data_source", "comparability_status",
        "proxy_quality", "methodology_note",
    }
    incomplete_benchmarks = [
        row["cell_id"] for row in applicable if any(not str(row.get(field, "")).strip() for field in benchmark_fields)
    ]
    check("benchmark_registry_104", not incomplete_benchmarks, "104/104 cells have a complete benchmark registry" if not incomplete_benchmarks else f"cells={', '.join(incomplete_benchmarks[:20])}")
    invalid_comparability = [row["cell_id"] for row in applicable if row.get("comparability_status") not in {"DIRECTLY_COMPARABLE", "LIMITED_COMPARABILITY", "STRUCTURAL_PROXY"}]
    check("comparability_status_valid", not invalid_comparability, "all comparability statuses are valid" if not invalid_comparability else f"cells={', '.join(invalid_comparability[:20])}")
    invalid_proxy = [row["cell_id"] for row in applicable if row.get("proxy_quality") not in {"DIRECT_REGIONAL", "DIRECT_GLOBAL", "REGIONAL_PROXY", "CROSS_REGIONAL_PROXY"}]
    check("no_invalid_or_missing_proxy", not invalid_proxy, "NO_VALID_PROXY=0 and all proxy qualities are valid" if not invalid_proxy else f"cells={', '.join(invalid_proxy[:20])}")
    cross_regional = [row["cell_id"] for row in applicable if row.get("proxy_quality") == "CROSS_REGIONAL_PROXY"]
    check("cross_regional_not_primary", not cross_regional, "no cell uses a cross-regional proxy as primary confirmation" if not cross_regional else f"cells={', '.join(cross_regional[:20])}")
    technical_bad = [row["cell_id"] for row in applicable if row.get("technical_status") != "TECHNICAL_PASS"]
    check("technical_status_104", not technical_bad, "TECHNICAL_PASS=104/104" if not technical_bad else f"cells={', '.join(technical_bad[:20])}")
    analytical_bad = [row["cell_id"] for row in applicable if row.get("analytical_quality") not in {"STRONG", "ACCEPTABLE"}]
    check("analytical_quality_no_review", not analytical_bad, "WEAK=0 and REVIEW_REQUIRED=0" if not analytical_bad else f"cells={', '.join(analytical_bad[:20])}")
    market_quality_bad = [row["cell_id"] for row in applicable if row.get("market_confirmation_quality") not in {"STRONG", "PARTIAL"}]
    check("market_confirmation_quality", not market_quality_bad, "WEAK=0 and CONTRADICTED=0" if not market_quality_bad else f"cells={', '.join(market_quality_bad[:20])}")
    evidence_by_cell: dict[str, list[dict[str, str]]] = {}
    for row in evidence:
        if row.get("review_status") == "PASS":
            evidence_by_cell.setdefault(row.get("cell_id", ""), []).append(row)
    unsupported = [row["cell_id"] for row in applicable if not evidence_by_cell.get(row["cell_id"])]
    check("evidence_package_per_cell", not unsupported, "104/104 cells have PASS evidence" if not unsupported else f"cells={', '.join(unsupported[:20])}")
    check("evidence_coverage_104", len(set(parsed_scores) & set(evidence_by_cell)) == 104, f"coverage={len(set(parsed_scores) & set(evidence_by_cell))}/104")

    direct_errors = []
    composite_errors = []
    rationale_errors = []
    cap_errors = []
    for score in applicable:
        rows = evidence_by_cell.get(score["cell_id"], [])
        roles = {row.get("component_role") for row in rows}
        if score["evidence_mode"] == "DIRECT":
            if not {"PRIMARY_RESEARCH", "MARKET_CONFIRMATION"} <= roles or not any(row.get("relevance_reason", "").startswith("DIRECT:") for row in rows):
                direct_errors.append(score["cell_id"])
            if abs(int(score["score"])) > 2:
                cap_errors.append(score["cell_id"])
        elif score["evidence_mode"] == "COMPOSITE":
            if not {"REGION_COMPONENT", "ASSET_COMPONENT", "MARKET_CONFIRMATION"} <= roles or len({row.get("source_id") for row in rows}) < 2:
                composite_errors.append(score["cell_id"])
            if any(not row.get("relevance_reason", "").strip() for row in rows):
                rationale_errors.append(score["cell_id"])
            if abs(int(score["score"])) > 1:
                cap_errors.append(score["cell_id"])
    check("direct_evidence_mapping", not direct_errors and any(row["evidence_mode"] == "DIRECT" for row in applicable), "all DIRECT mappings have primary research and market confirmation" if not direct_errors else f"cells={', '.join(direct_errors)}")
    check("composite_two_components", not composite_errors, "all COMPOSITE cells have region, asset and market components" if not composite_errors else f"cells={', '.join(composite_errors[:20])}")
    check("composite_relevance_rationale", not rationale_errors, "all COMPOSITE rationale fields are populated" if not rationale_errors else f"cells={', '.join(rationale_errors[:20])}")
    check("evidence_mode_score_caps", not cap_errors, "DIRECT and COMPOSITE caps passed" if not cap_errors else f"cells={', '.join(cap_errors)}")
    regional_market_errors = []
    for score in applicable:
        market_rows = [row for row in evidence_by_cell.get(score["cell_id"], []) if row.get("component_role") == "MARKET_CONFIRMATION"]
        if not market_rows or any(
            row.get("benchmark_id") != score.get("benchmark_id")
            or row.get("regional_scope") != score.get("geography")
            or row.get("URL") != score.get("market_data_source")
            for row in market_rows
        ):
            regional_market_errors.append(score["cell_id"])
    check("regional_market_confirmation_104", not regional_market_errors, "104/104 cells use the registered regional/global benchmark as market confirmation" if not regional_market_errors else f"cells={', '.join(regional_market_errors[:20])}")

    future_publications: list[str] = []
    future_market: list[str] = []
    stale_primary: list[str] = []
    invalid_urls: list[str] = []
    window_start = snapshot_date - timedelta(days=90)
    for row in evidence:
        try:
            publication = date.fromisoformat(row["publication_date"])
            market_date = date.fromisoformat(row["market_observation_date"])
        except (KeyError, ValueError):
            future_publications.append(row.get("evidence_id", "<missing date>"))
            continue
        if publication > snapshot_date:
            future_publications.append(row["evidence_id"])
        if market_date > snapshot_date:
            future_market.append(row["evidence_id"])
        if row.get("component_role") in {"PRIMARY_RESEARCH", "REGION_COMPONENT", "ASSET_COMPONENT"} and publication < window_start:
            stale_primary.append(row["evidence_id"])
        parsed = urlparse(row.get("URL", ""))
        domain = parsed.netloc.lower()
        if parsed.scheme not in {"http", "https"} or not domain or any(token in domain for token in ("example", "localhost", "demo", "fake")):
            invalid_urls.append(row.get("evidence_id", "<missing>"))
    check("publication_date_not_future", not future_publications, "all publication dates <= snapshot date" if not future_publications else f"evidence={', '.join(future_publications[:20])}")
    check("market_observation_not_future", not future_market, "all market observations <= snapshot date" if not future_market else f"evidence={', '.join(future_market[:20])}")
    check("primary_research_90_day_window", not stale_primary, "all primary research is inside the 90-day window" if not stale_primary else f"evidence={', '.join(stale_primary[:20])}")
    check("real_valid_urls", not invalid_urls, "all evidence URLs are valid and non-demo" if not invalid_urls else f"evidence={', '.join(invalid_urls[:20])}")
    window_days = int((metadata.get("data_windows") or {}).get("market_window_days") or 0)
    check("market_history_12_months", window_days >= 365, f"minimum market window={window_days} days")

    scenario_ids = {row.get("scenario_id") for row in scenarios}
    check("exactly_three_scenarios", len(scenarios) == 3 and scenario_ids == {"BASE", "UPSIDE", "DOWNSIDE"}, f"scenarios={len(scenarios)}")
    scenario_controls = all(row.get("trigger", "").strip() and row.get("veto", "").strip() for row in scenarios)
    check("scenario_trigger_veto", scenario_controls, "all scenarios have trigger and veto")
    affected_counts = [len([item for item in row.get("affected_cells", "").split(";") if item.strip()]) for row in scenarios]
    check("scenario_full_universe", affected_counts == [104, 104, 104], f"affected counts={affected_counts}")
    transmission_ids = [row.get("cell_id", "") for row in transmission]
    check("transmission_104", len(transmission) == 104 and len(set(transmission_ids)) == 104 and set(transmission_ids) == set(parsed_scores), f"rows={len(transmission)}, unique={len(set(transmission_ids))}")
    transmission_ok = all(row.get("control_status") == "PASS" and row.get("investment_logic", "").count("→") >= 3 for row in transmission)
    check("transmission_controls", transmission_ok, "104 driver → mechanism → asset → reaction chains PASS")
    threshold_bad = [
        row.get("cell_id", "") for row in transmission
        if not row.get("monitoring_indicator", "").strip()
        or not row.get("trigger_threshold", "").strip()
        or not row.get("veto_threshold", "").strip()
        or row.get("trigger_threshold") == row.get("veto_threshold")
        or not re.search(r"\d", row.get("trigger_threshold", ""))
        or not re.search(r"\d", row.get("veto_threshold", ""))
    ]
    check("quantified_distinct_trigger_veto", not threshold_bad, "104/104 monitoring controls are measurable and trigger differs from veto" if not threshold_bad else f"cells={', '.join(threshold_bad[:20])}")
    transmission_registry_bad = [
        row.get("cell_id", "") for row in transmission
        if row.get("benchmark_id") != full_scores.get(row.get("cell_id", ""), {}).get("benchmark_id")
        or row.get("comparability_status") != full_scores.get(row.get("cell_id", ""), {}).get("comparability_status")
    ]
    check("transmission_benchmark_synchronized", not transmission_registry_bad, "104/104 transmission rows use the score registry" if not transmission_registry_bad else f"cells={', '.join(transmission_registry_bad[:20])}")

    methodology_bad = []
    for row in applicable:
        if row.get("change_type") == "METHODOLOGY_REFINEMENT":
            try:
                if int(row["score_delta"]) == 0 or int(row["previous_score"]) + int(row["score_delta"]) != int(row["score"]):
                    methodology_bad.append(row["cell_id"])
            except (KeyError, TypeError, ValueError):
                methodology_bad.append(row["cell_id"])
    check("methodology_refinement_classified", not methodology_bad, "all methodology score changes reconcile and are not upgrades/downgrades" if not methodology_bad else f"cells={', '.join(methodology_bad[:20])}")
    quality_consistent = (
        quality_report.get("status") == "PASS"
        and quality_report.get("technical_pass_count") == 104
        and quality_report.get("regional_market_confirmation_coverage") == 104
        and quality_report.get("analytical_quality", {}).get("WEAK") == 0
        and quality_report.get("analytical_quality", {}).get("REVIEW_REQUIRED") == 0
        and str(quality_report.get("methodology_refinements")) in quality_markdown
    )
    check("quality_report_consistent", quality_consistent, "JSON and Markdown quality reports reconcile" if quality_consistent else "quality report mismatch or non-PASS")

    workbook = load_workbook(paths["workbook"], data_only=False, read_only=True)
    check("exactly_five_excel_sheets", workbook.sheetnames == EXPECTED_SHEETS, f"sheets={workbook.sheetnames}")
    excel_matrix = _read_full_score_matrix(workbook["Strategy_Sentiment_Map"])
    expected_matrix = {(row["asset_segment"], row["geography"]): ("N/A" if row["applicable"] == "false" else int(row["score"])) for row in scores}
    check("excel_csv_scores_match", excel_matrix == expected_matrix, f"Excel={len(excel_matrix)}, CSV={len(expected_matrix)}")
    excel_evidence = _read_table_rows(workbook["Evidence_Expectation"], "evidence_id")
    excel_evidence_ids = [str(row.get("evidence_id", "")) for row in excel_evidence]
    csv_evidence_ids = [row["evidence_id"] for row in evidence]
    check("excel_csv_evidence_match", excel_evidence_ids == csv_evidence_ids, f"Excel={len(excel_evidence_ids)}, CSV={len(csv_evidence_ids)}")
    excel_scenarios = _read_table_rows(workbook["Scenarios"], "scenario name")
    excel_scenario_map = {str(row.get("scenario name", "")): str(row.get("causal chain", "")) for row in excel_scenarios[:3]}
    csv_scenario_map = {row["scenario_name"]: row["causal_chain"] for row in scenarios}
    markdown_scenario_ok = all(row["scenario_name"] in report_text and row["causal_chain"] in report_text for row in scenarios)
    check("scenario_dataset_synchronized", excel_scenario_map == csv_scenario_map and markdown_scenario_ok, "Excel, CSV and Markdown use one scenario dataset")
    formula_errors = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if cell.data_type == "e" or (isinstance(value, str) and any(token in value for token in FORMULA_ERRORS)):
                    formula_errors.append(f"{worksheet.title}!{cell.coordinate}={value}")
    workbook.close()
    check("workbook_opens_no_formula_errors", not formula_errors, "workbook opens and contains no formula errors" if not formula_errors else "; ".join(formula_errors[:20]))

    pilot_hashes = metadata.get("pilot_hashes") or {}
    pilot_path_keys = {
        "mae_latest.xlsx": "workbook",
        "mae_latest_scores.csv": "scores",
        "mae_latest_evidence.csv": "evidence",
        "mae_latest_scenarios.csv": "scenarios",
        "mae_latest_report.md": "report",
    }
    pilot_changes = []
    for name, expected_hash in pilot_hashes.items():
        path = (pilot_output_paths or {}).get(pilot_path_keys.get(name, ""), root / "outputs" / name)
        actual_hash = hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else "missing"
        if actual_hash != expected_hash:
            pilot_changes.append(name)
    check("pilot_outputs_unchanged", not pilot_changes and len(pilot_hashes) == 5, "all five pilot latest artifacts are byte-identical" if not pilot_changes else f"changed={', '.join(pilot_changes)}")

    if require_database:
        for name, condition, detail in _full_database_checks(snapshot_date, metadata, scores, evidence, root=root):
            check(name, condition, detail)
    return _finish_full_validation(paths["validation"], metadata, checks, errors, write_report, snapshot_date, quality_report=quality_report)


def _read_full_score_matrix(worksheet) -> dict[tuple[str, str], int | str]:
    values = list(worksheet.iter_rows(values_only=True))
    header_index = None
    headers: list[str] = []
    for index, row in enumerate(values):
        normalized = [str(value or "").strip() for value in row]
        if "Asset segment" in normalized and set(FULL_GEOGRAPHIES) <= set(normalized):
            header_index = index
            headers = normalized
            break
    if header_index is None:
        return {}
    result: dict[tuple[str, str], int | str] = {}
    for row in values[header_index + 1 : header_index + 1 + len(FULL_ASSETS)]:
        asset = str(row[headers.index("Asset segment")] or "").strip()
        for geography in FULL_GEOGRAPHIES:
            value = row[headers.index(geography)]
            if value == "N/A":
                result[(asset, geography)] = "N/A"
            else:
                try:
                    result[(asset, geography)] = int(value)
                except (TypeError, ValueError):
                    continue
    return result


def _full_database_checks(
    snapshot_date: date,
    metadata: dict[str, Any],
    scores: list[dict[str, str]],
    evidence: list[dict[str, str]],
    *,
    root: Path,
) -> list[tuple[str, bool, str]]:
    database_path = get_settings().sqlite_path if root == ROOT_DIR else root / "data" / "mae.db"
    if not database_path.exists():
        return [("full_snapshot_persisted", False, "SQLite database is missing")]
    snapshot_ok = cells = applicable = not_applicable = supported = scenarios = 0
    full_evidence_count = 0
    with sqlite3.connect(f"file:{database_path}?mode=ro", uri=True) as connection:
        row = connection.execute(
            "SELECT id, run_metadata FROM mae_snapshots WHERE snapshot_date = ? AND is_demo = 0",
            (snapshot_date.isoformat(),),
        ).fetchone()
        if row:
            run_metadata = json.loads(row[1]) if isinstance(row[1], str) else dict(row[1] or {})
            snapshot_ok = int(run_metadata.get("full_product_version") == FULL_PRODUCT_VERSION)
            snapshot_id = row[0]
            cells = connection.execute("SELECT COUNT(*) FROM mae_snapshot_cells WHERE snapshot_id = ?", (snapshot_id,)).fetchone()[0]
            applicable = connection.execute("SELECT COUNT(*) FROM mae_snapshot_cells WHERE snapshot_id = ? AND applicability = 'APPLICABLE'", (snapshot_id,)).fetchone()[0]
            not_applicable = connection.execute("SELECT COUNT(*) FROM mae_snapshot_cells WHERE snapshot_id = ? AND applicability = 'NOT_APPLICABLE'", (snapshot_id,)).fetchone()[0]
            supported = connection.execute("SELECT COUNT(DISTINCT canonical_cell_id) FROM evidence_items WHERE snapshot_id = ? AND review_status = 'PASS'", (snapshot_id,)).fetchone()[0]
            non_anchor_ids = [row["cell_id"] for row in scores if row["applicable"] == "true" and row["evidence_mode"] != "ANCHOR_EXISTING"]
            placeholders = ",".join("?" for _ in non_anchor_ids)
            full_evidence_count = connection.execute(
                f"SELECT COUNT(*) FROM evidence_items WHERE snapshot_id = ? AND canonical_cell_id IN ({placeholders}) AND review_status = 'PASS'",
                (snapshot_id, *non_anchor_ids),
            ).fetchone()[0]
        scenarios = connection.execute("SELECT COUNT(*) FROM market_scenarios WHERE snapshot_date = ? AND is_demo = 0", (snapshot_date.isoformat(),)).fetchone()[0]
    expected_new_evidence = sum(row["evidence_mode"] != "ANCHOR_EXISTING" for row in evidence)
    preflight = (metadata.get("database_preflight") or {}).get("tables") or {}
    current = database_state(root=root).get("tables") or {}
    history_ok = all(current.get(table, 0) >= count for table, count in preflight.items())
    backups = list((root / "data" / "backups").glob("*.db"))
    return [
        ("sqlite_history_preserved", history_ok and bool(backups), f"backups={len(backups)}, preflight={preflight}, current={current}"),
        ("full_snapshot_persisted", snapshot_ok == 1 and cells == 114 and applicable == 104 and not_applicable == 10, f"snapshot={snapshot_ok}, cells={cells}, applicable={applicable}, N/A={not_applicable}"),
        ("sqlite_evidence_coverage", supported == 104, f"supported cells={supported}/104"),
        ("sqlite_idempotent_evidence", full_evidence_count == expected_new_evidence, f"new full evidence rows={full_evidence_count}, expected={expected_new_evidence}"),
        ("sqlite_three_scenarios", scenarios == 3, f"scenario rows={scenarios}"),
    ]


def _finish_full_validation(
    validation_path: Path,
    metadata: dict[str, Any],
    checks: list[dict[str, str]],
    errors: list[str],
    write_report: bool,
    snapshot_date: date,
    quality_report: dict[str, Any] | None = None,
) -> dict[str, Any]:
    status = "FAIL" if errors else "PASS"
    report = {
        "status": status,
        "snapshot_date": snapshot_date.isoformat(),
        "total_intersections": 114,
        "applicable_cells": 104,
        "not_applicable_cells": 10,
        "evidence_coverage": next((row["detail"] for row in checks if row["name"] == "evidence_coverage_104"), "not checked"),
        "checked_at": datetime.now(UTC).isoformat(),
        "checks": checks,
        "errors": errors,
        "metadata": metadata,
        "quality": quality_report or {},
        "warnings": (quality_report or {}).get("warnings", []),
    }
    if write_report:
        validation_path.parent.mkdir(parents=True, exist_ok=True)
        validation_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report
