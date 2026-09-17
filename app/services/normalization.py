from __future__ import annotations

import csv
import hashlib
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.config import get_settings


CANONICAL_REGIONS = ["Global", "US", "Europe", "UK", "Japan", "EM ex China"]
GENERAL_ASSESSMENT_RANGE = "A1:H22"

REGION_TO_TEMPLATE = {
    "GLOBAL": "Global",
    "US": "US",
    "U.S.": "US",
    "USA": "US",
    "UNITED STATES": "US",
    "EUROPE": "Europe",
    "UK": "UK",
    "UNITED KINGDOM": "UK",
    "JAPAN": "Japan",
    "EM EX CHINA": "EM ex China",
    "EM EX CH": "EM ex China",
    "EM": "EM ex China",
}

TEMPLATE_TO_DOMAIN_REGION = {
    "Global": "Global",
    "US": "US",
    "Europe": "Europe",
    "UK": "UK",
    "Japan": "Japan",
    "EM ex Ch": "EM ex China",
    "EM ex China": "EM ex China",
}


@dataclass(frozen=True)
class TemplateRow:
    row_index: int
    asset_class: str
    asset_group: str
    asset_segment: str
    template_row_key: str
    source_asset_class: str = ""
    source_label: str = ""


@dataclass(frozen=True)
class CanonicalCell:
    canonical_cell_id: str
    row_index: int
    source_column: str
    asset_class: str
    asset_group: str
    asset_segment: str
    source_asset_class: str
    source_label: str
    template_row_key: str
    region: str
    applicability: str
    baseline_score: int | None


def content_hash(text: str) -> str:
    return hashlib.sha256(normalize_space(text).encode("utf-8")).hexdigest()


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def normalize_institution(value: str) -> str:
    return normalize_space(value).casefold()


def canonicalize_url(url: str | None) -> str | None:
    if not url:
        return None
    parts = urlsplit(url.strip())
    scheme = parts.scheme.lower()
    netloc = parts.netloc.lower()
    query = [
        (k, v)
        for k, v in parse_qsl(parts.query, keep_blank_values=True)
        if not k.lower().startswith("utm_") and k.lower() not in {"fbclid", "gclid"}
    ]
    return urlunsplit((scheme, netloc, parts.path.rstrip("/") or "/", urlencode(query), ""))


def normalize_region(value: str) -> str:
    key = normalize_space(value).upper()
    return TEMPLATE_TO_DOMAIN_REGION.get(REGION_TO_TEMPLATE.get(key, value), value)


def template_region(value: str) -> str:
    key = normalize_space(value).upper()
    if key in REGION_TO_TEMPLATE:
        return REGION_TO_TEMPLATE[key]
    for template, domain in TEMPLATE_TO_DOMAIN_REGION.items():
        if value == domain:
            return template
    return value


def make_template_row_key(asset_class: str, asset_group: str, asset_segment: str) -> str:
    return "|".join([normalize_space(asset_class), normalize_space(asset_group), normalize_space(asset_segment)])


@lru_cache(maxsize=16)
def parse_template_rows(path: Path | None = None) -> list[TemplateRow]:
    path = path or get_settings().mae_template_xlsx
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        rows_by_key: dict[str, TemplateRow] = {}
        for cell in canonical_cell_registry(path, include_not_applicable=True):
            rows_by_key.setdefault(
                cell.template_row_key,
                TemplateRow(
                    row_index=cell.row_index,
                    asset_class=cell.asset_class,
                    asset_group=cell.asset_group,
                    asset_segment=cell.asset_segment,
                    template_row_key=cell.template_row_key,
                    source_asset_class=cell.source_asset_class,
                    source_label=cell.source_label,
                ),
            )
        return sorted(rows_by_key.values(), key=lambda row: row.row_index)
    rows: list[TemplateRow] = []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        all_rows = list(reader)
    for index, row in enumerate(all_rows[2:], start=3):
        padded = row + ["", "", ""]
        if not any(padded[:3]):
            continue
        rows.append(
            TemplateRow(
                row_index=index,
                asset_class=normalize_space(padded[0]),
                asset_group=normalize_space(padded[1]),
                asset_segment=normalize_space(padded[2]),
                template_row_key=make_template_row_key(padded[0], padded[1], padded[2]),
                source_asset_class=normalize_space(padded[0]),
                source_label=normalize_space(padded[2]),
            )
        )
    return rows


@lru_cache(maxsize=16)
def mae_result_template_row_keys() -> set[str]:
    return {row.template_row_key for row in parse_template_rows()}


@lru_cache(maxsize=32)
def canonical_cell_registry(path: Path | None = None, include_not_applicable: bool = True) -> list[CanonicalCell]:
    path = path or get_settings().mae_template_xlsx
    wb = load_workbook(path, data_only=True)
    ws = wb["ex"] if "ex" in wb.sheetnames else wb[wb.sheetnames[0]]
    header_regions = [normalize_space(ws.cell(2, col).value or "") for col in range(3, 9)]
    regions = [template_region(region) for region in header_regions]
    cells: list[CanonicalCell] = []
    current_asset = ""
    for row_index in range(3, 22):
        if ws[f"A{row_index}"].value:
            current_asset = normalize_space(str(ws[f"A{row_index}"].value))
        source_label = normalize_space(str(ws[f"B{row_index}"].value or ""))
        if not source_label:
            continue
        row_key = xlsx_general_assessment_row_key(current_asset, source_label)
        if row_key is None:
            continue
        asset_class, asset_group, asset_segment = row_key.split("|", 2)
        for offset, region in enumerate(regions, start=3):
            raw_value = ws.cell(row_index, offset).value
            applicable = raw_value not in (None, "")
            if not applicable and not include_not_applicable:
                continue
            baseline = valid_score(int(raw_value)) if applicable else None
            cells.append(
                CanonicalCell(
                    canonical_cell_id=canonical_cell_id(row_index, region),
                    row_index=row_index,
                    source_column=get_column_letter(offset),
                    asset_class=asset_class,
                    asset_group=asset_group,
                    asset_segment=asset_segment,
                    source_asset_class=current_asset,
                    source_label=source_label,
                    template_row_key=row_key,
                    region=region,
                    applicability="APPLICABLE" if applicable else "NOT_APPLICABLE",
                    baseline_score=baseline,
                )
            )
    return cells


@lru_cache(maxsize=16)
def applicable_canonical_cells(path: Path | None = None) -> list[CanonicalCell]:
    return canonical_cell_registry(path, include_not_applicable=False)


@lru_cache(maxsize=32)
def canonical_cell_map(path: Path | None = None, include_not_applicable: bool = True) -> dict[tuple[str, str], CanonicalCell]:
    return {
        (cell.template_row_key, template_region(cell.region)): cell
        for cell in canonical_cell_registry(path, include_not_applicable=include_not_applicable)
    }


def canonical_cell_for(template_row_key: str, region: str, include_not_applicable: bool = True) -> CanonicalCell | None:
    return canonical_cell_map(include_not_applicable=include_not_applicable).get((template_row_key, template_region(region)))


def is_applicable_canonical_cell(template_row_key: str, region: str) -> bool:
    cell = canonical_cell_for(template_row_key, region, include_not_applicable=True)
    return bool(cell and cell.applicability == "APPLICABLE")


def canonical_cell_id(row_index: int, region: str) -> str:
    return f"GA:R{row_index:02d}:{_region_code(region)}"


def canonical_cell_id_for(template_row_key: str, region: str) -> str | None:
    cell = canonical_cell_for(template_row_key, region, include_not_applicable=True)
    return cell.canonical_cell_id if cell else None


def map_to_canonical_row_key(asset_class: str, asset_group: str, asset_segment: str) -> str | None:
    direct = make_template_row_key(asset_class, asset_group, asset_segment)
    if direct in mae_result_template_row_keys():
        return direct
    asset = normalize_space(asset_class).upper()
    group = normalize_space(asset_group).casefold()
    segment = normalize_space(asset_segment)
    segment_key = segment.casefold()
    combined = f"{asset} {group} {segment_key}"
    if asset in {"GOLD", "COMMODITIES", "COMMODITY"}:
        if "other precious" in segment_key or "silver" in segment_key or "platinum" in segment_key:
            return "COMMODITIES|Commodities|Other precious metals"
        if "gold" in segment_key or "bullion" in combined:
            return "COMMODITIES|Commodities|Gold"
    if asset in {"EQUITY", "EQUITIES"}:
        if "wide market" in combined or "broad" in combined:
            return "EQUITY|Wide Market|Wide Market"
        if "breadth" in combined or "equal weight" in combined or "equal-weight" in combined:
            return "EQUITY|Other categories|Market Breadth (Equal Weight)"
        if "preferred" in combined:
            return "EQUITY|Other categories|Preferred"
        if "value" in combined or "valuation" in combined or "cheap" in combined:
            return "EQUITY|Other categories|Value"
        if "growth" in combined or "technology" in combined or "software" in combined or "semiconductor" in combined:
            return "EQUITY|Other categories|Growth"
        if "small cap" in combined or "small-cap" in combined or "smid" in combined:
            return "EQUITY|Other categories|Small Cap"
        return None
    if asset == "FIXED INCOME":
        if "gov" in group or "government" in combined or "sovereign" in combined:
            if "long" in segment_key:
                return "FIXED INCOME|GOV|Long Term"
            if "short" in segment_key:
                return "FIXED INCOME|GOV|Short Term"
            if "mid" in segment_key or "medium" in segment_key:
                return "FIXED INCOME|GOV|Mid Term"
        if "hy" in group or "high yield" in combined:
            if "long" in segment_key:
                return "FIXED INCOME|CORP HY|Long Term"
            if "short" in segment_key:
                return "FIXED INCOME|CORP HY|Short Term"
            if "mid" in segment_key or "medium" in segment_key:
                return "FIXED INCOME|CORP HY|Mid Term"
        if "ig" in group or "investment grade" in combined:
            if "long" in segment_key:
                return "FIXED INCOME|CORP IG|Long Term"
            if "short" in segment_key:
                return "FIXED INCOME|CORP IG|Short Term"
            if "mid" in segment_key or "medium" in segment_key:
                return "FIXED INCOME|CORP IG|Mid Term"
        if "inflation linked" in combined or "tips" in combined:
            return "FIXED INCOME|Other categories|Inflation Linked"
        if "mortgage" in combined or "mbs" in combined:
            return "FIXED INCOME|Other categories|Mortgage-Backed"
    return None


def xlsx_general_assessment_row_key(asset_class: str, label: str) -> str | None:
    asset = normalize_space(asset_class).upper()
    label_norm = normalize_space(label)
    if not label_norm:
        return None
    equity_other = {"Market Breadth (Equal Weight)", "Preferred", "Value", "Growth", "Small Cap"}
    if asset == "EQUITIES":
        if label_norm == "Wide Market":
            return "EQUITY|Wide Market|Wide Market"
        if label_norm in equity_other:
            return f"EQUITY|Other categories|{label_norm}"
        return None
    if asset == "FIXED INCOME":
        fixed_map = {
            "Government Bonds Long Term": "FIXED INCOME|GOV|Long Term",
            "Government Bonds Short Term": "FIXED INCOME|GOV|Short Term",
            "Government Bonds Mid Term": "FIXED INCOME|GOV|Mid Term",
            "Corporate IG Long Term": "FIXED INCOME|CORP IG|Long Term",
            "Corporate IG Short Term": "FIXED INCOME|CORP IG|Short Term",
            "Corporate IG Mid Term": "FIXED INCOME|CORP IG|Mid Term",
            "Corporate High Yield Long Term": "FIXED INCOME|CORP HY|Long Term",
            "Corporate High Yield Short Term": "FIXED INCOME|CORP HY|Short Term",
            "Corporate High Yield Mid Term": "FIXED INCOME|CORP HY|Mid Term",
            "Inflation Linked": "FIXED INCOME|Other categories|Inflation Linked",
            "Mortgage-Backed": "FIXED INCOME|Other categories|Mortgage-Backed",
        }
        return fixed_map.get(label_norm)
    if asset == "COMMODITIES":
        commodity_map = {
            "Gold": "COMMODITIES|Commodities|Gold",
            "Other precious metals": "COMMODITIES|Commodities|Other precious metals",
        }
        return commodity_map.get(label_norm)
    return None


def _region_code(region: str) -> str:
    return {
        "Global": "GLOBAL",
        "US": "US",
        "Europe": "EUROPE",
        "UK": "UK",
        "Japan": "JAPAN",
        "EM ex China": "EM_EX_CHINA",
        "EM ex Ch": "EM_EX_CHINA",
    }.get(template_region(region), re.sub(r"[^A-Z0-9]+", "_", region.upper()).strip("_"))


def valid_score(value: int) -> int:
    if value < -3 or value > 3:
        raise ValueError("score must be between -3 and +3")
    return value


def clip_score(value: float) -> float:
    return max(-3.0, min(3.0, value))


def round_half_away_from_zero(value: float) -> int:
    abs_value = abs(value)
    rounded = int(abs_value + 0.5)
    return rounded if value >= 0 else -rounded
